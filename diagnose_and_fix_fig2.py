"""Diagnóstico y reemplazo definitivo de la Figura 2 en Articulo_SP5_completo.docx"""
import os
import zipfile
import shutil
import docx

HERE = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(HERE, "Articulo_SP5_completo.docx")
FIG2_PATH = os.path.join(HERE, "backend", "article_outputs", "fig2_eda.png")

def inspect_and_fix():
    if not os.path.exists(DOCX_PATH):
        print("No existe DOCX_PATH")
        return

    # 1. Inspeccionar el archivo ZIP directamente
    with zipfile.ZipFile(DOCX_PATH, 'r') as z:
        names = z.namelist()
        media = [n for n in names if 'media' in n]
        print(f"\n[1] Archivos de imagen encontrados en word/media/:")
        for m in media:
            info = z.getinfo(m)
            print(f"    - {m} (tamaño: {info.file_size:,} bytes)")

    # 2. Cargar con docx para inspeccionar párrafos y tablas
    try:
        doc = docx.Document(DOCX_PATH)
    except PermissionError:
        print("\n[AVISO] Microsoft Word tiene abierto el archivo. Por favor, CIERRA Word y vuelve a ejecutar.\n")
        return

    print(f"\n[2] Búsqueda de menciones a 'Figura 2' en párrafos principales:")
    for idx, p in enumerate(doc.paragraphs):
        if "figura 2" in p.text.lower():
            print(f"    - Párrafo #{idx}: '{p.text[:90]}...'")

    print(f"\n[3] Búsqueda de menciones a 'Figura 2' en tablas:")
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    if "figura 2" in p.text.lower():
                        print(f"    - Tabla #{t_idx}, Fila #{r_idx}, Col #{c_idx}, Párrafo #{p_idx}: '{p.text[:90]}...'")

    # 4. Buscar qué imagen corresponde a la Figura 2 en el XML de document.xml
    with zipfile.ZipFile(DOCX_PATH, 'r') as z:
        xml_content = z.read('word/document.xml').decode('utf-8', errors='ignore')
        rels_content = z.read('word/_rels/document.xml.rels').decode('utf-8', errors='ignore')

    import xml.etree.ElementTree as ET
    rels_tree = ET.fromstring(rels_content)
    rid_to_target = {}
    for rel in rels_tree:
        rid = rel.attrib.get('Id')
        tgt = rel.attrib.get('Target')
        if rid and tgt:
            rid_to_target[rid] = tgt

    print(f"\n[4] Relaciones de imágenes en document.xml.rels:")
    for rid, tgt in rid_to_target.items():
        if "media" in tgt:
            print(f"    - {rid} -> {tgt}")

    # 5. Localizar el fragmento XML donde aparece "Figura 2" y ver qué imagen está antes
    pos_fig2 = xml_content.find("Figura 2")
    if pos_fig2 != -1:
        snippet_before = xml_content[max(0, pos_fig2 - 2000):pos_fig2]
        import re
        # Buscar r:embed="..." o r:id="..."
        embeds = re.findall(r'(?:embed|id)="([^"]+)"', snippet_before)
        print(f"\n[5] Referencias rId encontradas en los 2000 caracteres antes de 'Figura 2': {embeds}")
        
        target_media = None
        for eid in reversed(embeds):
            if eid in rid_to_target and "media" in rid_to_target[eid]:
                target_media = "word/" + rid_to_target[eid].lstrip('/')
                print(f"    -> [MATCH] La imagen asociada a Figura 2 es: {target_media}")
                break

        # Reemplazar la imagen en el archivo DOCX
        if target_media:
            with open(FIG2_PATH, 'rb') as f_img:
                new_img_data = f_img.read()

            # Leer todos los archivos del zip, reemplazar target_media, y guardar
            with zipfile.ZipFile(DOCX_PATH, 'r') as zin:
                all_files = {name: zin.read(name) for name in zin.namelist()}

            all_files[target_media] = new_img_data
            print(f"\n[6] Reemplazando {target_media} con nueva imagen de 4 paneles ({len(new_img_data):,} bytes)...")

            with zipfile.ZipFile(DOCX_PATH, 'w', zipfile.ZIP_DEFLATED) as zout:
                for name, data in all_files.items():
                    zout.writestr(name, data)
            print("[OK] Imagen reemplazada físicamente con éxito dentro del docx.")

            # Ahora actualizar el texto del pie de figura con python-docx
            doc = docx.Document(DOCX_PATH)
            new_caption_body = (
                "Análisis exploratorio de datos de la prevalencia de obesidad en adultos (CDC PLACES) y determinantes sociales en 1.632 tractos censales. "
                "(a) Distribución empírica con curva de densidad KDE, marcas de tendencia central (media = 29,92 %, mediana = 28,70 %, moda = 25,10 %) y umbrales por cuartiles (Q1 = 24,2 %, Q2 = 28,7 %, Q3 = 36,0 %). "
                "(b) Distribución territorial en Cook County, IL (n = 1.328). "
                "(c) Distribución territorial en New York County / Manhattan, NY (n = 304). "
                "(d) Matriz de correlación de rangos de Spearman entre determinantes sociales y obesidad."
            )

            # Buscar el pie de figura y actualizarlo
            for p in doc.paragraphs:
                if p.text.strip().startswith("Figura 2"):
                    p.text = ""
                    r_b = p.add_run("Figura 2. ")
                    r_b.bold = True
                    p.add_run(new_caption_body)
                    print("[OK] Pie de Figura 2 actualizado con formato.")
                    break

            for t in doc.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if p.text.strip().startswith("Figura 2"):
                                p.text = ""
                                r_b = p.add_run("Figura 2. ")
                                r_b.bold = True
                                p.add_run(new_caption_body)
                                print("[OK] Pie de Figura 2 actualizado dentro de tabla.")
                                break

            doc.save(DOCX_PATH)
            print(f"\n[ÉXITO DEFINITIVO] Archivo actualizado guardado en: {DOCX_PATH}")

if __name__ == "__main__":
    inspect_and_fix()
